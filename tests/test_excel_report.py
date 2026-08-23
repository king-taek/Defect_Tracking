"""「사진 대조」 Excel 리포트 구조 테스트 (design_handoff_fluent_redesign/04-excel-report.md).

Qt 없이 순수 로직만 검증한다. 실제 임시 이미지와 실제 썸네일 캐시를 써서 사진이 정말
만들어지고 셀 크기에 맞춰 줄어드는지까지 본다.
"""

from __future__ import annotations

import ast
import inspect
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.clustering import Cluster
from app.export import excel_report
from app.export.excel_report import (
    BLOCK_ROWS,
    EXPORT_WARN_BLOCKS,
    SHEET_TITLE,
    export_excel,
    photo_box_px,
)
from app.models import BaseDefectMatches, DefectRecord, MatchResult
from app.safety import OriginalProtectionError
from app.thumbnails import ThumbnailCache

LOT_LAYERS = ["LYA4", "LYB4", "LYA3", "LYB3"]
EMU_PER_PX = 9525


# ---- 픽스처 ----

@pytest.fixture()
def src(tmp_path) -> Path:
    """원본(LOT) 루트. 여기에는 절대 쓰지 않는다."""
    root = tmp_path / "204. DEVAINT.226 (PKG)"
    root.mkdir()
    return root


@pytest.fixture()
def cache(tmp_path) -> ThumbnailCache:
    return ThumbnailCache(tmp_path / "ws" / "cache")


def _photo(src: Path, name: str) -> Path:
    path = src / name
    if not path.exists():
        Image.new("RGB", (320, 320), (28, 38, 58)).save(path, format="JPEG", quality=80)
    return path


def _rec(src: Path, name: str, layer: str, *, col: int = 2, row: int = 4,
         x: float = 2210.0, y: float = 1180.0, defect: str = "PARTICLE") -> DefectRecord:
    return DefectRecord(
        image_path=_photo(src, name), wafer_id="W455685703", layer=layer,
        layer_folder=layer, col=col, row=row, x=x, y=y, defect_name=defect,
    )


def _item(src: Path, base_layer: str = "LYA4", *, tag: str = "a") -> BaseDefectMatches:
    """기준 1건 + LYB4 매칭 / LYA3 미매칭 / LYB3 매칭."""
    base = _rec(src, f"{tag}_2_4_2210_1180.jpg", base_layer)
    hit1 = _rec(src, f"{tag}_2_4_2214_1183.jpg", "LYB4")
    hit2 = _rec(src, f"{tag}_2_4_2207_1176.jpg", "LYB3")
    return BaseDefectMatches(base=base, results=[
        MatchResult(compare_layer="LYB4", base=base, matched=hit1, distance=5.0),
        # 미매칭이지만 진단용 최근접 거리는 들고 있다. 리포트에는 절대 나오면 안 된다.
        MatchResult(compare_layer="LYA3", base=base, matched=None,
                    nearest=hit2, nearest_distance=141.8, die_candidates=2),
        MatchResult(compare_layer="LYB3", base=base, matched=hit2, distance=8.0),
    ])


def _export(tmp_path: Path, src: Path, cache: ThumbnailCache,
            selected: list[BaseDefectMatches], **kwargs):
    out = kwargs.pop("out", tmp_path / "ws" / "exports" / "r.xlsx")
    params = dict(
        lot_name="204. DEVAINT.226 (PKG)", base_layer="LYA4",
        compare_layers=["LYB4", "LYA3", "LYB3"], tolerance=100.0,
        selected=selected, thumb_cache=cache, source_roots=[src],
        layer_order=LOT_LAYERS,
    )
    params.update(kwargs)
    return export_excel(out, **params)


def _sheet(path: Path):
    return load_workbook(path)[SHEET_TITLE]


def _values(ws) -> list[str]:
    return [
        str(cell.value)
        for row in ws.iter_rows() for cell in row
        if cell.value is not None
    ]


def _block_rows(ws) -> list[int]:
    """블록 머리 행 번호들."""
    return [
        r for r in range(1, ws.max_row + 1)
        if str(ws.cell(row=r, column=1).value or "").startswith("#")
    ]


def _rgb(color) -> str:
    value = getattr(color, "rgb", None)
    return value if isinstance(value, str) else ""


def _is_reddish(argb: str) -> bool:
    if len(argb) != 8:
        return False
    try:
        red, green, blue = (int(argb[i:i + 2], 16) for i in (2, 4, 6))
    except ValueError:
        return False
    return red > green + 40 and red > blue + 40


# ---- 시트 골격 ----

def test_single_sheet_named_photo_compare(tmp_path, src, cache):
    ws_path = _export(tmp_path, src, cache, [_item(src)])
    book = load_workbook(ws_path)
    assert book.sheetnames == ["사진 대조"]
    ws = book[SHEET_TITLE]
    assert ws.cell(row=1, column=1).value == "사진 대조"
    head = " ".join(str(ws.cell(row=1, column=c).value or "")
                    for c in range(1, ws.max_column + 1))
    assert "기준 LYA4" in head and "허용 오차 100 µm" in head
    assert "담은 1건" in head and "layer 4개" in head
    rule = " ".join(str(ws.cell(row=2, column=c).value or "")
                    for c in range(1, ws.max_column + 1))
    assert "열 = LOT layer 순서 고정" in rule and "★ 표시로만 구분" in rule
    assert "미매칭 = 회색 처리 · 거리 표기 없음" in rule


def test_freeze_panes_is_a3(tmp_path, src, cache):
    ws = _sheet(_export(tmp_path, src, cache, [_item(src)]))
    assert ws.freeze_panes == "A3"


def test_print_setup_is_landscape_fit_to_width(tmp_path, src, cache):
    ws = _sheet(_export(tmp_path, src, cache, [_item(src)]))
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert str(ws.print_title_rows).replace("$", "") == "1:2"


def test_one_page_break_per_block(tmp_path, src, cache):
    selected = [_item(src, tag=t) for t in ("a", "b", "c")]
    ws = _sheet(_export(tmp_path, src, cache, selected))
    heads = _block_rows(ws)
    assert len(heads) == 3
    breaks = [brk.id for brk in ws.row_breaks.brk]
    assert breaks == [head + BLOCK_ROWS - 1 for head in heads]


def test_block_is_five_rows_head_layer_photo_state_file(tmp_path, src, cache):
    """블록 = 머리 / layer / 사진 / 상태 / 파일명 5행. 사진 행은 96 px 상당."""
    assert BLOCK_ROWS == 5
    ws = _sheet(_export(tmp_path, src, cache, [_item(src)]))
    head = _block_rows(ws)[0]
    assert head == 3  # 머리 2행 바로 다음
    assert ws.cell(row=head + 1, column=1).value == "LYA4 ★ 기준"
    assert ws.cell(row=head + 2, column=1).value is None      # 사진 칸(이미지)
    assert ws.cell(row=head + 3, column=1).value == "기준"
    assert ws.cell(row=head + 4, column=1).value == "a_2_4_2210_1180.jpg"
    assert ws.row_dimensions[head + 2].height == pytest.approx(72.0)
    assert ws.max_row == head + BLOCK_ROWS - 1


# ---- 열 순서 ----

def test_column_order_follows_lot_layer_order(tmp_path, src, cache):
    """열은 LOT layer 순서. 기준(LYA4 가 아닌 LYB4)을 맨 앞으로 끌어오지 않는다."""
    base = _rec(src, "b_2_4_2210_1180.jpg", "LYB4")
    other = _rec(src, "b_2_4_2214_1183.jpg", "LYA4")
    item = BaseDefectMatches(base=base, results=[
        MatchResult(compare_layer="LYA4", base=base, matched=other, distance=4.0),
        MatchResult(compare_layer="LYA3", base=base, matched=None),
        MatchResult(compare_layer="LYB3", base=base, matched=None),
    ])
    ws = _sheet(_export(tmp_path, src, cache, [item], base_layer="LYB4"))
    layer_row = _block_rows(ws)[0] + 1
    labels = [ws.cell(row=layer_row, column=c).value for c in range(1, 5)]
    assert labels == ["LYA4", "LYB4 ★ 기준", "LYA3", "LYB3"]
    assert labels[0] != "LYB4 ★ 기준"  # 기준이 첫 열로 끌려나오지 않았다


def test_base_column_is_not_pulled_to_front_without_layer_order(tmp_path, src, cache):
    """layer_order 미전달 폴백에서도 기준은 제자리에 남는다."""
    base = _rec(src, "c_2_4_2210_1180.jpg", "L2")
    hit = _rec(src, "c_2_4_2214_1183.jpg", "L1")
    item = BaseDefectMatches(base=base, results=[
        MatchResult(compare_layer="L1", base=base, matched=hit, distance=4.0),
        MatchResult(compare_layer="L3", base=base, matched=None),
    ])
    ws = _sheet(_export(tmp_path, src, cache, [item], base_layer="L2",
                        compare_layers=["L1", "L3"], layer_order=None))
    layer_row = _block_rows(ws)[0] + 1
    assert [ws.cell(row=layer_row, column=c).value for c in (1, 2, 3)] == [
        "L1", "L2 ★ 기준", "L3"
    ]


def test_re_review_depth_suffix_in_layer_row(tmp_path, src, cache):
    base = _rec(src, "d_2_4_2210_1180.jpg", "LYA4")
    hit = _rec(src, "d_2_4_2214_1183.jpg", "LYA4_재재리뷰")
    item = BaseDefectMatches(base=base, results=[
        MatchResult(compare_layer="LYA4_재재리뷰", base=base, matched=hit, distance=3.0),
    ])
    ws = _sheet(_export(tmp_path, src, cache, [item],
                        compare_layers=["LYA4_재재리뷰"],
                        layer_order=["LYA4", "LYA4_재재리뷰"]))
    layer_row = _block_rows(ws)[0] + 1
    assert [ws.cell(row=layer_row, column=c).value for c in (1, 2)] == [
        "LYA4 ★ 기준", "LYA4 재재"
    ]


# ---- 미매칭은 회색 1단 ----

def test_unmatched_cell_is_grey_one_step(tmp_path, src, cache):
    ws = _sheet(_export(tmp_path, src, cache, [_item(src)]))
    head = _block_rows(ws)[0]
    col = 3  # LYA3 (미매칭)
    layer_cell = ws.cell(row=head + 1, column=col)
    photo_cell = ws.cell(row=head + 2, column=col)
    state_cell = ws.cell(row=head + 3, column=col)
    file_cell = ws.cell(row=head + 4, column=col)
    assert layer_cell.value == "LYA3"
    assert _rgb(layer_cell.fill.fgColor) == "FFEFF1F3"
    assert _rgb(layer_cell.font.color) == "FF8A949E"
    assert photo_cell.value == "미매칭"
    assert _rgb(photo_cell.fill.fgColor) == "FFF6F7F8"
    assert _rgb(photo_cell.font.color) == "FF8A949E"
    assert _rgb(state_cell.fill.fgColor) == "FFF6F7F8"
    assert file_cell.value is None  # 미매칭 칸에는 파일명이 없다
    # 사유·최근접 거리는 한 글자도 새지 않는다.
    joined = " ".join(_values(ws))
    for leak in ("허용오차 초과", "최근접", "141.8", "OVER_TOLERANCE"):
        assert leak not in joined


def test_no_red_fill_and_no_distance_text(tmp_path, src, cache):
    selected = [_item(src, tag=t) for t in ("a", "b")]
    ws = _sheet(_export(tmp_path, src, cache, selected))
    for row in ws.iter_rows():
        for cell in row:
            fill = _rgb(cell.fill.fgColor)
            font_color = _rgb(cell.font.color)
            assert fill != "FFB00020" and font_color != "FFB00020"  # 옛 _NOMATCH
            assert not _is_reddish(fill), (cell.coordinate, fill)
            assert not _is_reddish(font_color), (cell.coordinate, font_color)
    # 블록 안(3행 이후)에는 거리 표기가 없다. 머리 2행의 "허용 오차 100 µm" 는 설정값이다.
    body = " ".join(
        str(cell.value) for row in ws.iter_rows(min_row=3) for cell in row
        if cell.value is not None
    )
    assert "거리" not in body
    assert "µm" not in body
    assert not re.search(r"\d+\.\d", body)
    assert "매칭 없음" not in body and "매칭 O" not in body and "매칭 X" not in body


def test_match_state_is_only_two_words(tmp_path, src, cache):
    ws = _sheet(_export(tmp_path, src, cache, [_item(src)]))
    head = _block_rows(ws)[0]
    states = [ws.cell(row=head + 3, column=c).value for c in range(1, 5)]
    assert states == ["기준", "매칭", "-", "매칭"]
    matched_cell = ws.cell(row=head + 3, column=2)
    assert _rgb(matched_cell.font.color) == "FF1F7A1F"
    assert _rgb(matched_cell.fill.fgColor) == "FFEEF7EE"


# ---- 하이퍼링크 ----

def test_hyperlink_text_is_the_real_filename(tmp_path, src, cache):
    item = _item(src)
    ws = _sheet(_export(tmp_path, src, cache, [item]))
    links = {
        cell.hyperlink.target: cell
        for row in ws.iter_rows() for cell in row if cell.hyperlink is not None
    }
    expected = {
        Path(rec.image_path).resolve().as_uri(): Path(rec.image_path).name
        for rec in [item.base, item.results[0].matched, item.results[2].matched]
    }
    assert set(links) == set(expected)
    for target, cell in links.items():
        assert cell.value == expected[target]
        assert _rgb(cell.font.color) == "FF0F6CBD"
        assert cell.font.underline == "single"
    assert "원본 사진 열기" not in _values(ws)


# ---- 사진 ----

@pytest.mark.parametrize("scale", [1.0, 1.3])
def test_images_are_resized_to_the_cell_box(tmp_path, src, cache, scale):
    ws = _sheet(_export(tmp_path, src, cache, [_item(src)], font_scale=scale))
    box_w, box_h = photo_box_px(scale)
    assert ws._images, "사진이 한 장도 들어가지 않았다"
    for image in ws._images:
        ext = image.anchor.ext
        width, height = ext.cx / EMU_PER_PX, ext.cy / EMU_PER_PX
        assert 0 < width <= box_w and 0 < height <= box_h
        # 정사각 원본이므로 짧은 변에 정확히 맞아야 한다(앵커만 걸린 게 아니다).
        assert width == pytest.approx(min(box_w, box_h))
        assert height == pytest.approx(min(box_w, box_h))
    assert ws.row_dimensions[_block_rows(ws)[0] + 2].height == pytest.approx(72.0 * scale)


def test_cluster_extra_count_shows_plus_badge(tmp_path, src, cache):
    item = _item(src)
    extra = _rec(src, "a_2_4_2212_1182.jpg", "LYA4")
    item.base_cluster = Cluster(representative=item.base, members=[item.base, extra])
    ws = _sheet(_export(tmp_path, src, cache, [item]))
    head = _block_rows(ws)[0]
    badge = ws.cell(row=head + 2, column=1)
    assert badge.value == "＋1"
    assert _rgb(badge.fill.fgColor) == "FF000000"  # 사진 뒤 검정 바탕


def test_missing_photo_falls_back_to_text(tmp_path, src, cache):
    base = _rec(src, "e_2_4_2210_1180.jpg", "LYA4")
    gone = DefectRecord(image_path=src / "없는파일.jpg", wafer_id="W455685703",
                        layer="LYB4", layer_folder="LYB4", col=2, row=4, x=2210.0, y=1180.0)
    item = BaseDefectMatches(base=base, results=[
        MatchResult(compare_layer="LYB4", base=base, matched=gone, distance=2.0),
    ])
    ws = _sheet(_export(tmp_path, src, cache, [item], compare_layers=["LYB4"]))
    head = _block_rows(ws)[0]
    assert ws.cell(row=head + 2, column=2).value == "사진 없음"


# ---- 기준 없음(교차매치) ----

def test_mixed_base_layers_drop_the_star_badge(tmp_path, src, cache):
    first = _item(src, "LYA4", tag="a")
    second = _item(src, "LYB4", tag="b")
    second.results = [
        MatchResult(compare_layer="LYA4", base=second.base,
                    matched=second.results[0].matched, distance=6.0),
        MatchResult(compare_layer="LYA3", base=second.base, matched=None),
        MatchResult(compare_layer="LYB3", base=second.base,
                    matched=second.results[2].matched, distance=9.0),
    ]
    ws = _sheet(_export(tmp_path, src, cache, [first, second]))
    joined = " ".join(_values(ws))
    assert "★" not in joined
    # 블록 안에는 기준 배지도 "기준" 상태 칸도 없다(머리의 "기준 없음" 안내만 남는다).
    body = " ".join(
        str(cell.value) for row in ws.iter_rows(min_row=3) for cell in row
        if cell.value is not None
    )
    assert "기준" not in body
    head_text = " ".join(str(ws.cell(row=1, column=c).value or "")
                         for c in range(1, ws.max_column + 1))
    assert "기준 없음" in head_text and "조사 layer 4개" in head_text
    for row in _block_rows(ws):
        text = " ".join(
            str(ws.cell(row=row, column=c).value or "")
            for c in range(1, ws.max_column + 1)
        )
        assert "교차매치" in text
        # 기준 열이 없으므로 분모는 그 블록의 실제 열 수(기준 사진도 한 칸으로 센다).
        assert "매칭 3 / 4" in text
    layer_row = _block_rows(ws)[0] + 1
    assert [ws.cell(row=layer_row, column=c).value for c in range(1, 5)] == LOT_LAYERS
    assert _rgb(ws.cell(row=layer_row, column=1).fill.fgColor) == "FFE8EEF4"


# ---- 담은 것만 / 안전 게이트 / 상수 ----

def test_only_selected_entries_are_exported(tmp_path, src, cache):
    everything = [_item(src, tag=t) for t in ("a", "b", "c", "d")]
    selected = everything[:2]
    ws = _sheet(_export(tmp_path, src, cache, selected))
    assert len(_block_rows(ws)) == 2
    kept = {
        Path(rec.image_path).resolve().as_uri()
        for item in selected
        for rec in [item.base] + [mr.matched for mr in item.results if mr.matched]
    }
    targets = {
        cell.hyperlink.target
        for row in ws.iter_rows() for cell in row if cell.hyperlink is not None
    }
    assert targets <= kept
    dropped = {Path(everything[3].base.image_path).name}
    assert not (dropped & set(_values(ws)))


def test_export_into_source_root_is_blocked(tmp_path, src, cache):
    with pytest.raises(OriginalProtectionError):
        _export(tmp_path, src, cache, [_item(src)], out=src / "exports" / "leak.xlsx")
    assert not list(src.rglob("*.xlsx"))


def test_assert_output_safe_is_the_first_statement():
    """원본 보호 게이트가 함수 첫 문장이어야 한다(다른 작업보다 먼저 막는다)."""
    tree = ast.parse(inspect.getsource(export_excel))
    body = tree.body[0].body
    first = body[1] if isinstance(body[0], ast.Expr) else body[0]  # docstring 다음
    assert isinstance(first, ast.Assign)
    assert isinstance(first.value, ast.Call)
    assert getattr(first.value.func, "id", "") == "assert_output_safe"


def test_export_warn_blocks_constant():
    """분할은 하지 않고 UI 가 출력 전에 경고하도록 임계값만 노출한다(REVIEW-01 A13)."""
    assert EXPORT_WARN_BLOCKS == 200


def test_palette_constants_match_the_spec():
    assert excel_report._HEAD == "FF20303F"
    assert excel_report._SUBHEAD == "FFF1F5F9"
    assert excel_report._LYHEAD == "FFE8EEF4"
    assert excel_report._BLOCK == "FFEEF3F8"
    assert excel_report._GREY_BG == "FFF6F7F8"
    assert excel_report._GREY_FG == "FF8A949E"
    assert excel_report._PASS == "FF1F7A1F"
    assert excel_report._PASS_BG == "FFEEF7EE"
    assert excel_report._LINK == "FF0F6CBD"
    assert excel_report._LINE == "FFDCE3EA"
    # 걷어낸 옛 상수는 다시 들어오지 않는다.
    for gone in ("_NEON", "_NOMATCH", "_NAVY", "_LIGHT", "_GREY"):
        assert not hasattr(excel_report, gone)


def test_report_source_has_no_em_dash_or_emoji():
    """카피 규칙: em-dash 금지, 이모지 금지. 주석·docstring 에도 적용(REVIEW-01 AD5)."""
    em_dash = chr(0x2014)
    for path in (Path(excel_report.__file__), Path(__file__)):
        text = path.read_text(encoding="utf-8")
        assert em_dash not in text, path.name
        assert not [ch for ch in text if 0x1F300 <= ord(ch) <= 0x1FAFF], path.name
